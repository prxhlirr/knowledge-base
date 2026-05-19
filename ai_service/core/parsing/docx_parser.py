import os
import zipfile
import xml.etree.ElementTree as ET

from core.parsing.document_element import DocumentElement


def _docx_unique_cells(row):
    """
    合并单元格去重辅助函数：同一 _tc XML 节点只输出一次，消除重复列。
    根因：row.cells 对合并区域每个逻辑格均返回主格内容，产生重复列数据。
    """
    _seen, _cells = set(), []
    for _c in row.cells:
        _cid = id(_c._tc)
        if _cid not in _seen:
            _seen.add(_cid)
            _cells.append(_c.text.strip().replace('\n', ' '))
    return _cells


def _extract_sdt_text(sdt_elem) -> str:
    """
    辅助函数：从 w:sdt（Structured Document Tag）元素中递归提取纯文本。
    根因：公文模板常将机构名、文号等元数据封装在 sdt 控件中，
          python-docx 的标准段落遍历会跳过此类节点，导致内容静默丢失。
    """
    import re as _re
    # 提取所有 w:t 文本节点
    _ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    texts = [t.text or '' for t in sdt_elem.iter(f'{{{_ns}}}t')]
    return ''.join(texts).strip()


def _extract_header_lines(doc_obj) -> list:
    """
    业务功能：提取 Word 文档页眉（Header）中的所有文本行。

    根因：标准中国公文（GB/T 9704）模板常将红头区域（机构名称、文号、签发人行、
          红色分割线）全部放置在 Word 页眉（w:hdr）中，而非正文 body。
          python-docx 遍历 element.body 时完全无法访问页眉，导致文号/机构名
          从未出现在 raw_text 中，正则提取静默失败，document_number 永远为空。

    设计决策：
      - 提取段落文本 + 单行表格平铺（与 body 处理保持一致）
      - 将页眉内容前置在 body 内容之前，确保正则扫描 header_text[:1500] 时能命中
      - 跳过 is_linked_to_previous=True 的页眉（与前一节相同，无需重复）
      - 只取第一个非链接页眉（公文只有一个页眉）
    """
    _header_lines = []
    try:
        for _sec in doc_obj.sections:
            _hdr = _sec.header
            if _hdr is None or _hdr.is_linked_to_previous:
                continue
            # 提取页眉中的段落
            for _para in _hdr.paragraphs:
                _txt = _para.text.strip()
                if _txt:
                    _header_lines.append(_txt)
            # 提取页眉中的表格（文号/签发人单行双列布局）
            for _tbl in _hdr.tables:
                if not _tbl.rows:
                    continue
                if len(_tbl.rows) == 1:
                    # 单行表格：平铺输出保留文号原始格式
                    _cells = _docx_unique_cells(_tbl.rows[0])
                    _flat = '  '.join(c for c in _cells if c)
                    if _flat:
                        _header_lines.append(_flat)
                else:
                    # 多行表格：逐行平铺
                    for _row in _tbl.rows:
                        _cells = _docx_unique_cells(_row)
                        _flat = '  '.join(c for c in _cells if c)
                        if _flat:
                            _header_lines.append(_flat)
            # 只处理第一个有效页眉，避免多节文档重复
            if _header_lines:
                break
    except Exception as _e:
        print(f'⚠️ [docx_parser] 页眉提取异常（已忽略）: {_e}')
    return _header_lines


def _extract_docx_part_texts(file_path: str, part_name: str) -> list:
    """Extract plain text lines from a docx XML part such as comments/footnotes."""
    if not file_path:
        return []
    try:
        with zipfile.ZipFile(file_path) as zf:
            if part_name not in zf.namelist():
                return []
            root = ET.fromstring(zf.read(part_name))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        lines = []
        for node in root:
            texts = [t.text or "" for t in node.findall(".//w:t", ns)]
            text = "".join(texts).strip()
            if text:
                lines.append(text)
        return lines
    except Exception as err:
        print(f"⚠️ [docx_parser] 读取 {part_name} 失败，已跳过: {err}")
        return []


def _extract_header_elements(doc_obj) -> list:
    return [
        DocumentElement(text=line, type="header", source="docx", keep_for_index=False, keep_for_metadata=True)
        for line in _extract_header_lines(doc_obj)
    ]


def _extract_footer_elements(doc_obj) -> list:
    elements = []
    try:
        for sec in doc_obj.sections:
            footer = sec.footer
            if footer is None or footer.is_linked_to_previous:
                continue
            for para in footer.paragraphs:
                text = para.text.strip()
                if text:
                    elements.append(DocumentElement(text=text, type="footer", source="docx"))
            for tbl in footer.tables:
                for row in tbl.rows:
                    flat = "  ".join(c for c in _docx_unique_cells(row) if c)
                    if flat:
                        elements.append(DocumentElement(text=flat, type="footer", source="docx"))
    except Exception as err:
        print(f"⚠️ [docx_parser] 页脚提取异常（已忽略）: {err}")
    return elements


def _append_body_elements(doc_obj, lines: list, elements: list) -> None:
    for line in lines:
        stripped = (line or "").strip()
        if not stripped:
            continue
        typ = "title" if stripped.startswith("#") else "table" if stripped.startswith("|") else "paragraph"
        elements.append(DocumentElement(text=stripped, type=typ, source="docx"))


def _parse_docx_to_elements(doc_obj, file_path: str = None) -> list:
    elements = []
    elements.extend(_extract_header_elements(doc_obj))
    body_text = _parse_docx_to_markdown(doc_obj)
    header_lines = set(_extract_header_lines(doc_obj))
    body_lines = [line for line in body_text.splitlines() if line.strip() and line.strip() not in header_lines]
    _append_body_elements(doc_obj, body_lines, elements)
    elements.extend(_extract_footer_elements(doc_obj))
    if file_path:
        for text in _extract_docx_part_texts(file_path, "word/comments.xml"):
            elements.append(DocumentElement(text=text, type="comment", source="docx", keep_for_index=False))
        for text in _extract_docx_part_texts(file_path, "word/footnotes.xml"):
            elements.append(DocumentElement(text=text, type="footnote", source="docx", keep_for_index=False, keep_for_metadata=True))
        for text in _extract_docx_part_texts(file_path, "word/endnotes.xml"):
            elements.append(DocumentElement(text=text, type="endnote", source="docx", keep_for_index=False, keep_for_metadata=True))
    return elements


def _parse_docx_to_markdown(doc_obj) -> str:
    """
    业务功能：将 python-docx Document 对象转换为 Markdown 格式字符串。
    关键流程：按 XML body 元素顺序遍历（保留表格在段落中的原始位置，修复 BUG-04）。
    设计决策：提取为独立函数，供 DocxParser 和 DocParser(G路) 复用。
    [修复] 补充 sdt 节点提取，避免公文红头区域内容静默丢失。
    [修复] 单行双列表格（文号/签发人布局）直接平铺输出，保留文号原始格式供正则匹配。
    [修复] 提取 Word 页眉（Header）内容并前置插入，解决公文红头放在页眉导致文号完全丢失。
    """
    from docx.text.paragraph import Paragraph as _Para
    from docx.table import Table as _Table

    # [修复] 先提取页眉内容（公文红头常在此）
    _lines = _extract_header_lines(doc_obj)

    for _child in doc_obj.element.body:
        _ltag = _child.tag.split('}')[-1] if '}' in _child.tag else _child.tag
        if _ltag == 'p':
            _para = _Para(_child, doc_obj)
            _ptxt = _para.text.strip()
            if not _ptxt:
                continue
            _sname = (_para.style.name if _para.style else '').lower()
            if   'heading 1' in _sname or '标题 1' in _sname: _lines.append(f'# {_ptxt}')
            elif 'heading 2' in _sname or '标题 2' in _sname: _lines.append(f'## {_ptxt}')
            elif 'heading 3' in _sname or '标题 3' in _sname: _lines.append(f'### {_ptxt}')
            elif 'heading 4' in _sname or '标题 4' in _sname: _lines.append(f'#### {_ptxt}')
            else: _lines.append(_ptxt)
        elif _ltag == 'tbl':
            _tbl = _Table(_child, doc_obj)
            if not _tbl.rows:
                continue
            # [修复] 单行双列表格：公文中文号/签发人布局（1行×2格）
            # 此类表格若按 Markdown 表格格式输出，会在文号左右插入 | 符号，
            # 导致 _extract_dynamic_meta 中的文号正则无法从整行中独立匹配。
            # 直接平铺各格文本，保留原始内容，正则可正常命中文号格式。
            if len(_tbl.rows) == 1:
                _cells = _docx_unique_cells(_tbl.rows[0])
                _flat = '  '.join(c for c in _cells if c)
                if _flat:
                    _lines.append(_flat)
                continue
            # 多行表格：保持 Markdown 表格格式
            _lines.append('')
            _th = _docx_unique_cells(_tbl.rows[0])
            if _th:
                _lines.append('| ' + ' | '.join(_th) + ' |')
                _lines.append('| ' + ' | '.join(['---'] * len(_th)) + ' |')
            for _tr in _tbl.rows[1:]:
                _tc = _docx_unique_cells(_tr)
                if any(_tc):
                    _lines.append('| ' + ' | '.join(_tc) + ' |')
            _lines.append('')
        elif _ltag == 'sdt':
            # [修复] 提取 w:sdt 结构化文档标签内的文本
            # 公文模板常将机构名称、文号等封装在 sdt 控件中，标准遍历会静默跳过
            _sdt_text = _extract_sdt_text(_child)
            if _sdt_text:
                _lines.append(_sdt_text)

    return '\n'.join(_lines).strip()



class DocxParser:
    """
    业务功能：解析标准 .docx 文件，通过 python-docx 的 Heading 样式识别实现结构感知解析。
    核心价值：MarkItDown 无法识别 Word 原生 Heading 1/2/3 样式，所有段落被平铺为普通文本，
              导致 SemanticChunker 无法识别标题层次，切片质量大幅下降。
              python-docx 通过 paragraph.style.name 精准识别，输出真正的 # ## ### Markdown 标题。
    """

    def can_parse(self, file_path: str) -> bool:
        return file_path.lower().endswith(".docx")

    def parse(self, file_path: str) -> str:
        """
        解析 .docx 文件，按 XML body 顺序遍历段落和表格，输出 Markdown 格式文本。
        失败时返回空字符串，由 ParserFactory 降级至 MarkItDownParser。
        """
        try:
            import docx as _docx_mod
            _doc_obj = _docx_mod.Document(file_path)
            result = _parse_docx_to_markdown(_doc_obj)
            if result:
                _tbl_cnt = len(_doc_obj.tables)
                print(f'✅ [DocxParser] 结构化解析成功 ({len(result)} 字符, 含 {_tbl_cnt} 个表格, 位置已保留)')
                return result
            print('⚠️ [DocxParser] 提取内容为空，降级 MarkItDown')
            return ""
        except Exception as _err:
            print(f'⚠️ [DocxParser] 解析失败，降级 MarkItDown: {_err}')
            return ""

    def parse_elements(self, file_path: str) -> list:
        try:
            import docx as _docx_mod
            doc_obj = _docx_mod.Document(file_path)
            elements = _parse_docx_to_elements(doc_obj, file_path)
            print(f'[DocxParser] 结构化元素解析成功 ({len(elements)} elements)')
            return elements
        except Exception as err:
            print(f'⚠️ [DocxParser] 结构化解析失败，回退纯文本包装: {err}')
            text = self.parse(file_path)
            return [DocumentElement(text=line.strip(), type="paragraph", source="docx") for line in text.splitlines() if line.strip()]
